"""Offline security regression tests: synthetic accounts/data, no production env."""
import io
import os
import re
import sqlite3
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch, MagicMock
from urllib.parse import parse_qs, urlsplit

IMPORT_DIR = tempfile.TemporaryDirectory(prefix='cau-security-import-')
with patch.dict(os.environ, {
    'APP_API_KEY': 'test-api-key-not-for-deployment',
    'MINDMATE_PEPPER': 'test-pepper-not-for-deployment',
    'SECRET_KEY': 'test-session-secret-not-for-deployment-12345',
    'SUPABASE_URL': '', 'SUPABASE_KEY': '',
    'PORTAL_DB_PATH': str(Path(IMPORT_DIR.name) / 'portal.db'),
    'DATA_DIR': str(Path(IMPORT_DIR.name) / 'files'),
    'SEED_ADMIN_EMAIL': '', 'SEED_ADMIN_PASSWORD': '',
    'RATELIMIT_STORAGE_URI': 'memory://',
}), patch('dotenv.load_dotenv'), patch('requests.sessions.Session.request', side_effect=AssertionError('Network forbidden')):
    import app as portal

from security_utils import SafeCSVWriter, validate_archive, validate_image, neutralize_excel_formulas, validate_production_settings


class SecurityTests(unittest.TestCase):
    def test_production_never_uses_local_account_fallback(self):
        for missing in ('SUPABASE_URL', 'SUPABASE_KEY'):
            env = {'RENDER': 'true', 'SUPABASE_URL': 'https://example.test', 'SUPABASE_KEY': 'test'}
            env.pop(missing)
            with self.assertRaises(SystemExit):
                validate_production_settings(env)

    def test_production_cannot_disable_secure_cookie(self):
        with self.assertRaises(SystemExit):
            validate_production_settings({'RENDER': 'true', 'SUPABASE_URL': 'https://example.test',
                                          'SUPABASE_KEY': 'test', 'SESSION_COOKIE_SECURE': 'false'})
        validate_production_settings({'RENDER': 'true', 'SUPABASE_URL': 'https://example.test',
                                      'SUPABASE_KEY': 'test'})
        validate_production_settings({'SESSION_COOKIE_SECURE': 'false'})

    @classmethod
    def setUpClass(cls):
        cls.password_hash = portal.generate_password_hash('A-long-test-password!')

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix='cau-security-test-')
        self.addCleanup(self.temp.cleanup)
        portal.DB_PATH = str(Path(self.temp.name) / 'test.db')
        portal.DATA_DIR = portal.PORTAL_FILES_DIR = str(Path(self.temp.name) / 'files')
        portal.GALLERY_FOLDER = str(Path(self.temp.name) / 'gallery')
        portal.GALLERY_DATA_FILE = str(Path(self.temp.name) / 'gallery.json')
        portal.ADMIN_EMAILS = {'admin@example.test'}
        portal.app.config.update(TESTING=True, WTF_CSRF_ENABLED=True)
        portal.limiter.enabled = False
        with patch.dict(os.environ, {'SEED_ADMIN_EMAIL': '', 'SEED_ADMIN_PASSWORD': ''}):
            portal.init_db()
        self.network = patch('requests.sessions.Session.request', side_effect=AssertionError('Network forbidden'))
        self.network.start()
        self.addCleanup(self.network.stop)
        for email in ('owner@example.test', 'other@example.test', 'viewer@example.test', 'editor@example.test', 'admin@example.test'):
            portal._set_account(email, self.password_hash, already_hashed=True)
        for pid, email in [('own', 'owner@example.test'), ('other', 'other@example.test')]:
            portal.sb('POST', 'projects', data={'id': pid, 'name': pid, 'researcher_email': email})
            portal.sb('POST', 'project_participants', data={'id': pid+'-p', 'project_id': pid, 'code': pid+'-code'})
            portal.sb('POST', 'measurements', data={'id': pid+'-m', 'project_id': pid, 'participant_id': pid+'-p', 'data': {}})
            portal.sb('POST', 'project_variables', data={'id': pid+'-v', 'project_id': pid, 'name': 'score'})
            portal.sb('POST', 'tasks', data={'id': pid+'-t', 'project_id': pid, 'title': pid+'-task', 'created_by': email})
            portal.sb('POST', 'wardy_events', data={'id': pid+'-w', 'project_id': pid, 'user_name': pid+'-user', 'state': 'open'})
        for role in ('viewer', 'editor'):
            portal.sb('POST', 'project_collaborators', data={'project_id': 'own', 'researcher_email': role+'@example.test', 'role': role})
        self.client = portal.app.test_client()
        self.sign_in()

    def sign_in(self, email='owner@example.test'):
        sid = str(portal._uuid.uuid4())
        now = int(portal.time.time())
        portal.sb('POST', 'portal_auth_sessions', data={'id': sid, 'researcher_email': email, 'created_at': now, 'last_seen': now})
        with self.client.session_transaction() as sess:
            sess.clear()
            sess['researcher'] = email
            sess['auth_stamp'] = portal._auth_stamp(portal._get_account(email))
            sess['auth_session_id'] = sid

    def test_logout_rejects_replayed_cookie(self):
        self.token()
        cookie = self.client.get_cookie('session').value
        self.assertEqual(self.post('/logout').status_code, 302)
        self.client.set_cookie('session', cookie)
        self.assertEqual(self.client.get('/portal').status_code, 302)

    def test_tampered_cookie_rejected(self):
        cookie = self.client.get_cookie('session').value
        self.client.set_cookie('session', 'tampered'+cookie)
        self.assertEqual(self.client.get('/portal').status_code, 302)

    def test_session_store_failure_is_not_authorization(self):
        with sqlite3.connect(portal.DB_PATH) as conn:
            conn.execute('DROP TABLE portal_auth_sessions')
        self.assertEqual(self.client.get('/portal').status_code, 503)

    def test_contact_rate_limit(self):
        portal.limiter.enabled = True
        portal.limiter.reset()
        try:
            statuses = [self.post('/contact', {'message': 'test'}).status_code for _ in range(6)]
            self.assertEqual(statuses[:5], [200]*5)
            self.assertEqual(statuses[-1], 429)
        finally:
            portal.limiter.enabled = False

    def test_login_rate_limit(self):
        portal.limiter.enabled = True
        portal.limiter.reset()
        try:
            statuses = [self.post('/login', {'email': 'missing@example.test', 'password': 'wrong'}).status_code for _ in range(11)]
            self.assertEqual(statuses[-1], 429)
        finally:
            portal.limiter.enabled = False

    def test_redirect_does_not_accept_external_target(self):
        result = self.post('/login?next=https://attacker.invalid', {'email': 'owner@example.test', 'password': 'A-long-test-password!'})
        self.assertEqual(result.headers['Location'], '/portal')

    def test_mindmate_batch_is_bounded(self):
        for data in ('[]', '{"dataCount":999999999}', '{"dataCount":"bad"}'):
            response = self.client.post('/api/mindmate/followup', json={
                'order': 'updateData_daily_json', 'PhoneNumber': '01000000000', 'jsonData': data,
            }, headers={'X-API-Key': portal.APP_API_KEY})
            self.assertEqual(response.status_code, 400)

    def test_idle_and_absolute_session_timeout(self):
        for field, age in [('last_seen', 1801), ('created_at', 28801)]:
            self.sign_in()
            with self.client.session_transaction() as sess:
                sid = sess['auth_session_id']
            portal.sb('PATCH', 'portal_auth_sessions', data={field: int(portal.time.time())-age}, params='?id=eq.'+sid)
            self.assertEqual(self.client.get('/portal').status_code, 302)

    def token(self):
        page = self.client.get('/login').get_data(as_text=True)
        return re.search(r"csrfMeta.content = '([^']+)'", page)[1]

    def post(self, url, data=None, json=None):
        return self.client.post(url, data=data, json=json, headers={'X-CSRFToken': self.token()})

    def test_login_rotates_session_and_has_secure_cookie(self):
        with self.client.session_transaction() as sess:
            sess['old_marker'] = 'discard'
        response = self.post('/login', {'email': 'owner@example.test', 'password': 'A-long-test-password!'})
        self.assertEqual(response.status_code, 302)
        cookie = response.headers['Set-Cookie']
        for value in ('Secure', 'HttpOnly', 'SameSite=Lax'):
            self.assertIn(value, cookie)
        with self.client.session_transaction() as sess:
            self.assertNotIn('old_marker', sess)
            self.assertIn('auth_stamp', sess)

    def test_forged_old_session_rejected(self):
        with self.client.session_transaction() as sess:
            sess.pop('auth_stamp')
        self.assertEqual(self.client.get('/portal').status_code, 302)

    def test_password_change_revokes_existing_session(self):
        portal.sb('PATCH', 'accounts', data={'password': 'changed-hash'}, params='?email=eq.owner%40example.test')
        self.assertEqual(self.client.get('/portal').status_code, 302)

    def test_password_recovery_does_not_issue_token(self):
        response = self.post('/reset-password', {'email': 'other@example.test'})
        self.assertEqual(response.status_code, 302)
        self.assertEqual(portal.sb('GET', 'reset_tokens'), [])
        self.assertEqual(self.client.get('/reset-password/previously-issued-token').status_code, 404)

    def test_csrf_on_json_mutations_and_logout(self):
        for url in ('/portal/tasks/own-t/status', '/portal/projects/own/data/cell', '/logout'):
            self.assertEqual(self.client.post(url, json={'status': 'done'}).status_code, 400)
        self.assertEqual(self.client.get('/logout').status_code, 405)
        self.assertEqual(self.post('/logout').status_code, 302)

    def test_own_mutations_still_work(self):
        self.assertEqual(self.post('/portal/tasks/own-t/status', json={'status': 'done'}).status_code, 200)
        self.assertEqual(self.post('/portal/projects/own/data/cell', json={'measurement_id': 'own-m', 'column': 'score', 'value': 7}).status_code, 200)
        self.assertEqual(portal.sb('GET', 'measurements', params='?id=eq.own-m')[0]['data']['score'], 7)

    def test_foreign_tasks_cannot_be_changed(self):
        for suffix in ('status', 'edit', 'delete'):
            response = self.post('/portal/tasks/other-t/'+suffix, json={'status': 'done'})
            self.assertEqual(response.status_code, 403)
        self.assertEqual(portal.sb('GET', 'tasks', params='?id=eq.other-t')[0]['status'], 'todo')

    def test_nested_ids_cannot_cross_projects(self):
        for route in ('participants/other-p/edit', 'participants/other-p/delete',
                      'measurements/other-m/edit', 'measurements/other-m/delete',
                      'variables/other-v/edit', 'variables/other-v/delete'):
            self.assertEqual(self.post('/portal/projects/own/'+route).status_code, 404, route)
        self.assertEqual(self.client.get('/portal/projects/own/participants/other-p').status_code, 404)

    def test_hidden_participant_and_file_ids_checked(self):
        self.assertEqual(self.post('/portal/projects/own/measurements/add', {'participant_id': 'other-p'}).status_code, 404)
        portal.sb('POST', 'portal_files', data={'id': 'other-f', 'project_id': 'other', 'filename': 'other.csv', 'original_name': 'other.csv', 'researcher_email': 'other@example.test'})
        self.assertEqual(self.client.get('/portal/projects/own/data?file_id=other-f').status_code, 404)

    def test_viewer_read_allowed_write_denied(self):
        self.sign_in('viewer@example.test')
        self.assertEqual(self.client.get('/portal/projects/own').status_code, 200)
        self.assertEqual(self.post('/portal/projects/own/edit', {'name': 'changed'}).status_code, 403)

    def test_editor_cannot_delete_project_or_app_rows(self):
        self.sign_in('editor@example.test')
        self.assertEqual(self.post('/portal/projects/own/delete').status_code, 403)
        self.post('/portal/wardy/delete', {'ids': 'own-w'})
        self.assertTrue(portal.sb('GET', 'wardy_events', params='?id=eq.own-w'))

    def test_merge_requires_source_ownership(self):
        self.assertEqual(self.post('/portal/merge/preview', {'project1_id': 'own', 'project2_id': 'other'}).status_code, 403)
        self.assertEqual(self.post('/portal/merge/execute', {'p1_id': 'own', 'p2_id': 'other'}).status_code, 403)
        self.assertEqual(len(portal.sb('GET', 'projects')), 2)

    def test_unscoped_views_do_not_leak_other_projects(self):
        for url in ('/portal', '/portal/wardy', '/portal/wardy/export', '/portal/all/participants'):
            response = self.client.get(url)
            self.assertEqual(response.status_code, 200, url)
            self.assertNotIn('other-user', response.get_data(as_text=True))
            self.assertNotIn('other-code', response.get_data(as_text=True))
        for url in ('/portal/wardy', '/portal/neurobreeze'):
            self.assertEqual(self.client.get(url+'?project_id=other').status_code, 403)

    def test_collaborator_can_see_scoped_activity(self):
        self.sign_in('viewer@example.test')
        body = self.client.get('/portal/wardy').get_data(as_text=True)
        self.assertIn('own-user', body)
        self.assertNotIn('other-user', body)

    def test_admin_actions_fail_closed(self):
        for url in ('/portal/contacts', '/portal/audit', '/portal/settings/backup', '/portal/drive', '/portal/mindmate', '/portal/mindmate/live/data'):
            self.assertEqual(self.client.get(url).status_code, 403, url)
        self.assertEqual(self.post('/portal/settings/add-account', {'email': 'evil@example.test', 'password': 'very-long-password'}).status_code, 403)
        self.sign_in('admin@example.test')
        self.assertEqual(self.post('/portal/settings/restore').status_code, 403)

    def test_query_values_cannot_inject_filters(self):
        bad = 'owner@example.test&email=neq.x#fragment'
        with patch.object(portal, '_USE_SUPABASE', True), patch.object(portal._requests, 'get') as get:
            get.return_value.json.return_value = []
            portal._get_account(bad)
            query = parse_qs(urlsplit(get.call_args.args[0]).query)
            self.assertEqual(query['email'], ['eq.'+bad])
            self.assertEqual(set(query), {'email', 'select'})
        self.assertIsNone(portal._get_account(bad))

    def test_sql_values_parameterized_and_columns_validated(self):
        self.assertEqual(portal.sb('GET', 'projects', params='?id=eq.'+portal._q("' OR 1=1 --")), [])
        with self.assertRaises(ValueError):
            portal.sb('POST', 'projects', data={'name) VALUES (1); --': 'x'})

    def test_api_key_not_rendered(self):
        for url in ('/portal/wardy', '/portal/neurobreeze', '/api/docs'):
            self.assertNotIn(portal.APP_API_KEY, self.client.get(url).get_data(as_text=True))

    def test_api_auth_and_invalid_payload(self):
        url = '/api/wardy/events'
        self.assertEqual(self.client.post(url, json={}).status_code, 401)
        for payload in ([], [None], 'string', [1], [{'state': 'open'}]*501):
            self.assertEqual(self.client.post(url, json=payload, headers={'X-API-Key': portal.APP_API_KEY}).status_code, 400)
        response = self.client.post(url, json={'state': 'open', 'user_name': 'test', 'project_id': 'own'}, headers={'X-API-Key': portal.APP_API_KEY})
        self.assertEqual(response.status_code, 201)

    def test_no_default_supabase_seed(self):
        with patch.object(portal, '_USE_SUPABASE', True), patch.dict(os.environ, {'PORTAL_EMAIL': '', 'PORTAL_PASSWORD': ''}), patch.object(portal, '_set_account') as create:
            portal._seed_supabase_account()
            create.assert_not_called()

    def test_headers_methods_and_unnecessary_files(self):
        response = self.client.get('/portal')
        self.assertEqual(response.headers['Cache-Control'], 'no-store')
        self.assertEqual(response.headers['X-Content-Type-Options'], 'nosniff')
        self.assertEqual(response.headers['X-Frame-Options'], 'DENY')
        for path in ('/.env', '/.git/config', '/portal.db', '/app.py', '/static/', '/static/../app.py'):
            self.assertEqual(self.client.get(path).status_code, 404, path)
        for method in ('TRACE', 'PUT', 'DELETE'):
            self.assertEqual(self.client.open('/', method=method).status_code, 405)

    def test_path_traversal_and_symlinks_blocked(self):
        root = Path(self.temp.name) / 'root'
        root.mkdir()
        (root/'escape').symlink_to('/private/tmp')
        with portal.app.test_request_context():
            for name in ('../private', '/etc/passwd', 'escape/anything'):
                with self.assertRaises(portal.HTTPException):
                    portal._stored_path(str(root), name)

    def test_upload_executable_and_fake_image_rejected(self):
        self.assertEqual(self.post('/portal/files/upload', {'file': (io.BytesIO(b'evil'), 'evil.html')}).status_code, 400)
        self.sign_in('admin@example.test')
        with patch.object(portal, 'load_events', return_value=[{'key': 'test'}]):
            self.assertEqual(self.post('/gallery/test/upload', {'photo': (io.BytesIO(b'<script>alert(1)</script>'), 'image.png')}).status_code, 400)

    def test_valid_image_and_archive_bomb(self):
        from PIL import Image
        image = io.BytesIO()
        Image.new('RGB', (2, 2)).save(image, format='PNG')
        image.seek(0)
        validate_image(image, '.png')
        archive = io.BytesIO()
        with zipfile.ZipFile(archive, 'w', zipfile.ZIP_DEFLATED) as z:
            z.writestr('large.xml', b'A'*1_000_000)
        archive.seek(0)
        with self.assertRaises(ValueError):
            validate_archive(archive)

    def test_normal_spreadsheet_import_remains_usable(self):
        from openpyxl import Workbook
        workbook = Workbook()
        workbook.active.append(['code', 'score'])
        workbook.active.append(['new-code', 7])
        stream = io.BytesIO()
        workbook.save(stream)
        stream.seek(0)
        response = self.post('/portal/projects/own/upload', {'file': (stream, 'normal.xlsx')})
        self.assertEqual(response.status_code, 302)
        files = portal.sb('GET', 'portal_files', params='?project_id=eq.own')
        self.assertEqual(len(files), 1)
        import json
        self.assertEqual(len(json.loads(files[0]['column_schema'])), 2)

    def test_api_error_body_has_no_internal_details(self):
        with patch.object(portal, '_USE_SUPABASE', True), patch.object(portal._requests, 'get', side_effect=RuntimeError('PRIVATE-SERVER-URL')):
            response = self.client.get('/portal')
            self.assertNotIn('PRIVATE-SERVER-URL', response.get_data(as_text=True))

    def test_csv_and_excel_formulas_are_text(self):
        output = io.StringIO()
        SafeCSVWriter(output).writerow(['=HYPERLINK("bad")', '@SUM(1)', -2, 'ordinary'])
        self.assertIn("'=HYPERLINK", output.getvalue())
        self.assertIn("'@SUM", output.getvalue())
        from openpyxl import Workbook, load_workbook
        workbook = Workbook()
        workbook.active.append(['=1+1', 4])
        neutralize_excel_formulas(workbook)
        buf = io.BytesIO()
        workbook.save(buf)
        buf.seek(0)
        self.assertEqual(load_workbook(buf).active['A1'].data_type, 's')

    def test_all_templates_compile(self):
        for name in portal.app.jinja_env.list_templates():
            portal.app.jinja_env.get_template(name)

    def test_public_app_description_is_escaped(self):
        with patch.object(portal, 'load_extra_apps', return_value=[]), patch.object(portal, 'APPS', [{**portal.APPS[0], 'desc_kr': '<script>alert(1)</script>'}]):
            body = self.client.get('/apps/'+portal.APPS[0]['key']).get_data(as_text=True)
            self.assertNotIn('<script>alert(1)</script>', body)
            self.assertIn('&lt;script&gt;', body)

    def test_drive_outside_configured_folders_rejected(self):
        svc = MagicMock()
        svc.files.return_value.get.return_value.execute.return_value = {'name': 'private', 'mimeType': 'text/plain', 'parents': ['unrelated']}
        with patch.object(portal.drive_sync, '_service', return_value=svc):
            with self.assertRaises(PermissionError):
                portal.drive_sync.get_file('id', ['allowed'])
        svc.files.return_value.get_media.assert_not_called()


if __name__ == '__main__':
    unittest.main()
